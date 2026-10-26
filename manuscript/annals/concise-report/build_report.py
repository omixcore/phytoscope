"""Build the one-main-figure Concise Report from reviewed Markdown sources.
Use the Codex primary Python runtime for document generation.
"""
from pathlib import Path
from copy import deepcopy
import importlib.util,re,shutil,subprocess,xml.etree.ElementTree as E,json
from docx import Document
from docx.shared import Pt,Mm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
O=Path(__file__).resolve().parent;BASE=O.parent
spec=importlib.util.spec_from_file_location('phyto_docx',BASE/'build_docx.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
def base_doc():
 d=Document();s=d.sections[0];s.page_width=Mm(210);s.page_height=Mm(297)
 s.left_margin=s.right_margin=s.top_margin=s.bottom_margin=Mm(25)
 for name,size,bold in [('Normal',11,False),('Title',14,False),('Heading 1',11,True),('Heading 2',11,False)]:
  st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor(0,0,0)
  rf=st._element.get_or_add_rPr().rFonts
  for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+k),None)
  for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'Arial')
  st.paragraph_format.space_after=Pt(4)
  if 'Heading' in name:st.paragraph_format.keep_with_next=True
 d.styles['Normal'].paragraph_format.line_spacing=2
 # No footer or decorative title borders.
 for name in ['Title','Heading 1','Heading 2']:
  pp=d.styles[name]._element.get_or_add_pPr();bb=pp.find(qn('w:pBdr'))
  if bb is not None:pp.remove(bb)
 return d
def paragraph(d,text,style=None,spacing=None,size=None):
 p=d.add_paragraph(style=style);m.add_runs(p,text)
 if spacing:p.paragraph_format.line_spacing=spacing
 if size:
  for r in p.runs:r.font.name='Arial';r.font.size=Pt(size)
 return p
def make(src,dest,kind):
 d=base_doc();in_code=False;code=[];table=[];section=''
 if kind!='main':d.styles['Normal'].paragraph_format.line_spacing=1.15
 def flush():
  nonlocal table
  if not table:return
  data=[]
  for l in table:
   cells=[s.strip() for s in l.strip().strip('|').split('|')]
   if all(re.fullmatch(r':?-{3,}:?',s) for s in cells):continue
   data.append(cells)
  t=d.add_table(rows=0,cols=len(data[0]));t.style='Table Grid'
  for i,row in enumerate(data):
   cc=t.add_row().cells
   for c,s in zip(cc,row):
    c.text='';m.add_runs(c.paragraphs[0],s)
    if i==0:
     sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F0F2F3');c._tc.get_or_add_tcPr().append(sh)
   if i==0:
    tr=t.rows[0]._tr.get_or_add_trPr();hh=OxmlElement('w:tblHeader');hh.set(qn('w:val'),'true');tr.append(hh)
  table=[]
 for l in src.read_text().splitlines():
  if l.startswith('```'):
   flush()
   if in_code:
    p=paragraph(d,'\n'.join(code),spacing=1.0,size=8)
    for r in p.runs:r.font.name='Courier New'
    code=[];in_code=False
   else:in_code=True
   continue
  if in_code:code.append(l);continue
  if l.startswith('|'):table.append(l);continue
  flush()
  if not l.strip():continue
  if l.startswith('# '):paragraph(d,l[2:],style='Title',spacing=1.1)
  elif l.startswith('## '):
   section=l[3:];p=paragraph(d,section,style='Heading 1',spacing=1.0)
   if kind=='main' and section=='Introduction':p.paragraph_format.page_break_before=True
   if section.startswith('Table S'):
    for r in p.runs:r.font.size=Pt(7);r.bold=False
  elif l.startswith('### '):paragraph(d,l[4:],style='Heading 2',spacing=1.0)
  else:
   if section=='References':
    p=paragraph(d,l,spacing=1.1,size=10);p.paragraph_format.keep_together=True
   elif section=='Figure Legend':
    p=paragraph(d,l,spacing=1.15,size=7)
    for r in p.runs:r.bold=False
   else:
    p=paragraph(d,l)
    if kind=='main' and (l.startswith('**') or section=='Abstract'):
     p.paragraph_format.line_spacing=1.15
    if section=='Abstract':p.paragraph_format.keep_together=True
 flush();m.apply_publication_table_style(d);d.save(dest);return d
main=make(O/'manuscript.md',O/'manuscript.docx','main')
si=make(O/'supporting-information.md',O/'supporting-information.docx','support')
make(O/'cover-letter.md',O/'cover-letter.docx','cover')
# Keep the approved main-figure geometry. Uppercase labels follow Annals policy.
NS='http://www.w3.org/2000/svg';E.register_namespace('',NS)
r=E.parse(BASE/'figures/standardized/Figure_1.svg').getroot()
for t in r.iter('{'+NS+'}text'):
 if t.text in ['a','b','c','d']:t.text=t.text.upper()
p=O/'Figure_1.svg';E.ElementTree(r).write(p,encoding='unicode',xml_declaration=True)
for ext in ['pdf','png']:
 cmd=['inkscape',str(p),f'--export-type={ext}',f'--export-filename={O/f"Figure_1.{ext}"}']
 if ext=='png':cmd+=['--export-dpi=600']
 subprocess.run(cmd,check=True,capture_output=True)
subprocess.run(['inkscape',str(p),'--export-type=svg','--export-text-to-path','--export-plain-svg',f'--export-filename={O/"Figure_1_outlined.svg"}'],check=True,capture_output=True)
subprocess.run(['python',str(O/'consolidate_support.py')],check=True)
