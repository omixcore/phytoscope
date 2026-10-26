from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.enum.text import WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).parent

def set_style(style, size, bold=False):
    style.font.name = 'Times New Roman'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = __import__('docx').shared.RGBColor(0,0,0)

def add_runs(p, text):
    # Limited Markdown emphasis conversion, sufficient for the source draft.
    pat = re.compile(r'(\*\*.*?\*\*|\*[^*]+\*|`[^`]+`)')
    pos = 0
    for m in pat.finditer(text):
        if m.start() > pos: p.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith('**'):
            p.add_run(token[2:-2]).bold = True
        elif token.startswith('`'):
            run=p.add_run(token[1:-1]);run.font.name='Courier New';run.font.size=Pt(10)
        else:
            p.add_run(token[1:-1]).italic = True
        pos = m.end()
    if pos < len(text): p.add_run(text[pos:])

def apply_publication_table_style(doc):
    # Figure/table style requested by the authors; body text keeps manuscript styles.
    for table in doc.tables:
        table.autofit = False
        tblpr = table._tbl.tblPr
        indent = tblpr.find(qn('w:tblInd'))
        if indent is None:
            indent = OxmlElement('w:tblInd'); tblpr.append(indent)
        indent.set(qn('w:w'),'114'); indent.set(qn('w:type'),'dxa')
        widths = [int(c.width or Mm(40)) for c in table.columns]
        factor = int(Mm(159.5)) / sum(widths)
        widths = [int(w * factor) for w in widths]
        for col, width in zip(table.columns, widths):
            col.width = width
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell.width = widths[min(i, len(widths)-1)]
                tcpr = cell._tc.get_or_add_tcPr()
                borders = tcpr.find(qn('w:tcBorders'))
                if borders is None:
                    borders = OxmlElement('w:tcBorders'); tcpr.append(borders)
                for edge in ['top','left','bottom','right','insideH','insideV']:
                    el = borders.find(qn('w:'+edge))
                    if el is None:
                        el = OxmlElement('w:'+edge); borders.append(el)
                    for k,v in [('val','single'),('sz','4'),('color','000000'),('space','0')]:
                        el.set(qn('w:'+k),v)
                for para in cell.paragraphs:
                    para.paragraph_format.line_spacing = 1.0
                    for run in para.runs:
                        run.font.name = 'Arial'; run.font.size = Pt(7)
                        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial')
                        run.font.color.rgb = RGBColor(0,0,0)
                        run.bold = False; run.italic = False
    for para in doc.paragraphs:
        if re.match(r'^Table S[0-9]+[.]',para.text):
            for run in para.runs:
                run.font.name = 'Arial'; run.font.size = Pt(7)
                run.font.color.rgb = RGBColor(0,0,0)
                run.bold = False; run.italic = False

def parse_md(src, out, title_style=True):
    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=Inches(1); sec.bottom_margin=Inches(1); sec.left_margin=Inches(1); sec.right_margin=Inches(1)
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    for nm,size,bold in [('Normal',12,False),('Title',16,True),('Heading 1',14,True),('Heading 2',12,True),('Heading 3',12,True)]:
        set_style(doc.styles[nm],size,bold)
    doc.styles['Normal'].paragraph_format.line_spacing=2
    doc.styles['Normal'].paragraph_format.space_after=Pt(0)
    # Remove Word's built-in decorative Title paragraph border.
    title_ppr=doc.styles['Title']._element.get_or_add_pPr()
    title_border=title_ppr.find(qn('w:pBdr'))
    if title_border is not None: title_ppr.remove(title_border)
    for nm in ['Heading 1','Heading 2','Heading 3']:
        doc.styles[nm].paragraph_format.keep_with_next=True
        doc.styles[nm].paragraph_format.space_before=Pt(12)
        doc.styles[nm].paragraph_format.space_after=Pt(4)
        doc.styles[nm].paragraph_format.line_spacing=1.3
    lines=src.read_text().splitlines()
    in_code=False; code=[]; table=[]
    def flush_table():
        nonlocal table
        if not table:return
        parsed=[]
        for line in table:
            cells=[c.strip() for c in line.strip().strip('|').split('|')]
            if all(re.fullmatch(r':?-{3,}:?',c) for c in cells): continue
            parsed.append(cells)
        if parsed:
            t=doc.add_table(rows=0, cols=len(parsed[0]))
            t.style='Table Grid'
            for ri,row in enumerate(parsed):
                cells=t.add_row().cells
                for ci,val in enumerate(row):
                    cells[ci].text=''
                    p=cells[ci].paragraphs[0]; p.paragraph_format.line_spacing=1.05
                    p.paragraph_format.space_after=Pt(2)
                    add_runs(p,val)
                    for r in p.runs:
                        r.font.name='Times New Roman'; r.font.size=Pt(9)
                        if ri==0:r.bold=True
                    if ri==0:
                        for cell in [cells[ci]]:
                            tcPr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'D9E2F3');tcPr.append(shd)
            # header row repeat flag
            trPr=t.rows[0]._tr.get_or_add_trPr(); hdr=OxmlElement('w:tblHeader');hdr.set(qn('w:val'),'true');trPr.append(hdr)
        table=[]
    for line in lines:
        if line.startswith('```'):
            flush_table()
            if in_code:
                p=doc.add_paragraph();p.paragraph_format.line_spacing=1.0
                r=p.add_run('\n'.join(code));r.font.name='Courier New';r.font.size=Pt(9)
                code=[];in_code=False
            else:in_code=True
            continue
        if in_code:code.append(line);continue
        if line.startswith('|'):
            table.append(line);continue
        flush_table()
        if not line.strip():continue
        if line.startswith('# '):
            p=doc.add_paragraph(style='Title' if title_style else 'Heading 1');add_runs(p,line[2:].strip())
        elif line.startswith('## '):
            p=doc.add_paragraph(style='Heading 1');add_runs(p,line[3:].strip())
            if line[3:].strip() == 'Abstract': p.paragraph_format.page_break_before=True
        elif line.startswith('### '):
            p=doc.add_paragraph(style='Heading 2');add_runs(p,line[4:].strip())
        elif line.startswith('#### '):
            p=doc.add_paragraph(style='Heading 3');add_runs(p,line[5:].strip())
        elif line.startswith('- [ ] ') or line.startswith('- [x] '):
            p=doc.add_paragraph(style='List Bullet');add_runs(p,('☐ ' if '[ ]' in line[:7] else '☒ ')+line[7:])
        elif line.startswith('- '):
            p=doc.add_paragraph(style='List Bullet');add_runs(p,line[2:])
        else:
            p=doc.add_paragraph();add_runs(p,line)
            if line.startswith('**Figure '): p.paragraph_format.keep_together=True
    flush_table()
    # page-number footer
    for s in doc.sections:
        p=s.footer.paragraphs[0];p.alignment=2
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
    if src.stem == 'supplementary-information':
        for section in doc.sections:
            section.page_width = Mm(210); section.page_height = Mm(297)
            section.left_margin = section.right_margin = Mm(25)
            section.top_margin = section.bottom_margin = Mm(25)
    apply_publication_table_style(doc)
    doc.save(out)

if __name__ == '__main__':
    parse_md(BASE/'manuscript.md', BASE/'manuscript.docx')
    parse_md(BASE/'supplementary-information.md', BASE/'supplementary-information.docx')
